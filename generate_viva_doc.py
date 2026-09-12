import os
import sys
import subprocess

# Ensure python-docx is installed
try:
    import docx
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("python-docx is not installed. Installing it now...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "--break-system-packages"])
    import docx
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_docx():
    print("Generating CareerHub_Viva_Preparation_Guide.docx...")
    doc = docx.Document()

    # Define color scheme (Sleek Slate & Royal Blue theme)
    PRIMARY_COLOR = RGBColor(15, 23, 42)    # Slate 900
    SECONDARY_COLOR = RGBColor(37, 99, 235)  # Blue 600
    TEXT_COLOR = RGBColor(51, 65, 85)       # Slate 700
    DARK_TEXT = RGBColor(15, 23, 42)        # Slate 900

    # Adjust default normal style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = TEXT_COLOR

    # --- Title Section ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("CAREERHUB PORTAL\n")
    title_run.font.size = Pt(26)
    title_run.font.bold = True
    title_run.font.color.rgb = PRIMARY_COLOR

    subtitle_run = title_p.add_run("End-Term Presentation & Viva Examination Preparation Guide")
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.italic = True
    subtitle_run.font.color.rgb = SECONDARY_COLOR

    doc.add_paragraph().paragraph_format.space_after = Pt(24)

    # Helper function for custom headings
    def add_custom_heading(text, level, space_before=18, space_after=6):
        h = doc.add_heading(text, level=level)
        h.paragraph_format.space_before = Pt(space_before)
        h.paragraph_format.space_after = Pt(space_after)
        h.paragraph_format.keep_with_next = True
        run = h.runs[0]
        run.font.name = 'Calibri'
        if level == 1:
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = PRIMARY_COLOR
        elif level == 2:
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = SECONDARY_COLOR
        elif level == 3:
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = DARK_TEXT
        return h

    # --- SECTION 1: INTRODUCTION & OVERVIEW ---
    add_custom_heading("1. Project Overview & Architecture", level=1)
    
    p1 = doc.add_paragraph()
    p1.add_run("CareerHub ").bold = True
    p1.add_run("is a modern, production-grade MERN-stack application designed to bridge the gap between Indian freshers (candidates), corporate recruiters (employers), and college placement administrators. It solves traditional campus hiring pain points by combining structured profile building, direct job boards, and automated AI assistance.")

    add_custom_heading("Technical Stack Summary", level=2)
    
    # Table of Tech Stack
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Light Shading Accent 1'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Layer'
    hdr_cells[1].text = 'Technology Used'
    hdr_cells[2].text = 'Key Purpose / Library'
    
    # Bold the headers
    for cell in hdr_cells:
        for p in cell.paragraphs:
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = PRIMARY_COLOR

    tech_data = [
        ("Frontend Client", "React, Vite, Tailwind CSS", "Single Page Application (SPA), Component-based, mobile-first design."),
        ("Routing", "React Router DOM v6", "Config-driven layout-based routing with programmatic paths."),
        ("Backend Server", "Node.js, Express.js", "RESTful API handling business rules, request validation, and middlewares."),
        ("Database Layer", "MongoDB Atlas, Mongoose", "Document-oriented cloud database, schemas, relational validation, text indexes."),
        ("AI Integrations", "Google Gemini AI Node SDK", "Generative AI API for Mock Interview feedback and SSE-based chatbot."),
        ("Security Suite", "JWT, bcryptjs, Helmet, Cookie-Parser", "HTTP-only cookie sessions, password salting, CORS configurations, rate-limiters.")
    ]

    for layer, tech, purpose in tech_data:
        row_cells = table.add_row().cells
        row_cells[0].text = layer
        row_cells[1].text = tech
        row_cells[2].text = purpose

    # Space after table
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(12)

    # --- SECTION 2: VIVA PRESENTATION SCRIPT ---
    add_custom_heading("2. Viva Presentation & Live Demo Script", level=1)
    
    p2 = doc.add_paragraph()
    p2.add_run("This presentation script is structured for a ").italic = True
    p2.add_run("5 to 10-minute live demonstration").bold = True
    p2.add_run(" before the external examiner. Use it to speak confidently, outline decisions, and explain execution flow during your presentation.")

    add_custom_heading("Slide 1 & 2: Introduction & Project Scope (1 Minute)", level=2)
    p_slide1 = doc.add_paragraph()
    p_slide1.add_run("What to Say: ").bold = True
    p_slide1.add_run('"Good morning/afternoon respected examiners. Today I am presenting my final year project, CareerHub Portal. The core objective of this project is to streamline fresher hiring in Indian colleges. Traditional portals lack interactive support, so we built CareerHub as a decoupled client-server system that connects candidates, employers, and system admins. It provides features like a Live Resume Builder, a smart job matching portal, a real-time Career Chatbot, and an automated AI Mock Interview Simulator."')

    add_custom_heading("Slide 3 & 4: Architecture & Security (1.5 Minutes)", level=2)
    p_slide2 = doc.add_paragraph()
    p_slide2.add_run("What to Say: ").bold = True
    p_slide2.add_run('"Our system utilizes a decoupled architecture. The frontend is built on React and Vite for fast performance, styled with Tailwind CSS, and uses React Router DOM. The backend is an Express Node API connecting to MongoDB Atlas via Mongoose. For security, we do not store JWT credentials in local storage to prevent XSS attacks; instead, we implement HttpOnly cookies. We also use security headers with Helmet, request rate-limiting, and input sanitization to prevent query injections and cross-site scripting."')

    add_custom_heading("Live Demo Walkthrough (4-5 Minutes)", level=2)
    
    demo_p1 = doc.add_paragraph()
    demo_p1.add_run("Step 1: Admin Dashboard & Settings").bold = True
    demo_p1.paragraph_format.left_indent = Pt(18)
    demo_p1.add_run("\n• Action: Log in at '/admin/login' with credentials admin@careerhub.example | Admin@1234.\n• Show: Point out the stats dashboard showing the total registered recruiters, candidates, jobs, and placement rates.\n• Talk Track: \"Here is the Admin dashboard. The admin has the master controls to monitor user registration, delete fraudulent listings, and review active stats, providing a centralized control panel for college placement officers.\"")

    demo_p2 = doc.add_paragraph()
    demo_p2.add_run("Step 2: Candidate Live Resume Builder").bold = True
    demo_p2.paragraph_format.left_indent = Pt(18)
    demo_p2.add_run("\n• Action: Log in as Candidate (arjun@example.com) and navigate to Resume Builder.\n• Show: Type inside any input on the left (e.g. adding a new skill or project) and show how the PDF preview sheet on the right updates instantly. Click 'Save' to save it.\n• Talk Track: \"This is our Live Resume Builder. It leverages React's state binding. As the student types on the left panel, the preview canvas on the right updates in real-time. When the student clicks save, the backend commits this structured document to MongoDB under the candidate's record, which then automatically updates their searchable recruiter profile.\"")

    demo_p3 = doc.add_paragraph()
    demo_p3.add_run("Step 3: Recruiter Search Portal & Job Posting").bold = True
    demo_p3.paragraph_format.left_indent = Pt(18)
    demo_p3.add_run("\n• Action: Log in as Recruiter (hr@nimbus.example) and go to the Resumes Board.\n• Show: Type a skill in the search box (e.g., 'React' or 'Node') and press Enter. Show the filtered candidates. Then post a new job vacancy.\n• Talk Track: \"From the Recruiter portal, employers can post new jobs and search resumes. To power this, we created a MongoDB compound text index on candidate skills and job titles, allowing recruiters to perform high-speed keyword matching across hundreds of profiles instantly.\"")

    demo_p4 = doc.add_paragraph()
    demo_p4.add_run("Step 4: AI Mock Interview Simulator & Chatbot").bold = True
    demo_p4.paragraph_format.left_indent = Pt(18)
    demo_p4.add_run("\n• Action: Go back to the Candidate Dashboard, open 'Mock Interview', start a session, submit answers, and wait for the AI report card. Open the floating chat bubble and send a question.\n• Show: The AI grading panel with dynamic marks, feedback, and model answers. Show the bot responses streaming word-by-word.\n• Talk Track: \"Lastly, we integrated Google Gemini AI. The Candidate can start a role-specific mock interview. The system prompt instructs Gemini to output structured JSON with questions, evaluations, scores, and explanations. The feedback is stored as an InterviewSession. Meanwhile, our CareerBot uses Server-Sent Events (SSE) to stream real-time career advice over a single, persistent HTTP connection, avoiding the overhead of WebSockets.\"")


    # --- SECTION 3: VIVA QUESTIONS ---
    add_custom_heading("3. Comprehensive Viva Q&A (Top 26 Questions)", level=1)
    
    questions = [
        # PART A
        ("Q1: What is the overall architecture of this project?",
         "The project is structured as a decoupled full-stack Web Application utilizing the MERN stack (MongoDB, Express, React, Node.js). "
         "The frontend client is a React Single Page Application (SPA) compiled using Vite and styled with Tailwind CSS, running on port 5173. "
         "The backend is a RESTful API server powered by Node.js and Express, running on port 5001. "
         "The database is hosted on the cloud via MongoDB Atlas and accessed using Mongoose ODM. "
         "AI models are integrated on the backend via the Google Gemini AI SDK."),
         
        ("Q2: Why did you separate the frontend and backend into two separate directories?",
         "Separating frontend and backend represents a standard enterprise software pattern called Separation of Concerns (SoC). "
         "It decouples user interface (UI) rendering from data persistence and business logic. "
         "Benefits include: independent scalability (e.g. hosting the React build on Vercel/Netlify while deploying the Node server to Render/AWS), "
         "allowing frontend and backend developers to work in parallel without git conflicts, and "
         "making the codebase easier to test, maintain, and upgrade."),

        ("Q3: How does CORS work in this project, and why did you configure credentials?",
         "Cross-Origin Resource Sharing (CORS) is a browser security mechanism that restricts resources from being loaded by origins other than the server's own. "
         "Since the frontend (port 5173) and backend (port 5001) run on different origins, standard requests would be blocked. "
         "We resolved this by using the 'cors' middleware on the backend, whitelisting the frontend domain, and setting credentials: true. "
         "On the frontend, we set withCredentials: true on Axios instances. This allows secure HttpOnly session cookies (JWT) to pass across origins."),

        # PART B
        ("Q4: How is routing managed in your React application?",
         "We use React Router DOM v6. Instead of hardcoding paths across components, we created a centralized route configuration file "
         "(routerpath.jsx) mapping path definitions to their respective React page components. "
         "We defined a global ROUTES object containing constants like USER_DASHBOARD, ADMIN_LOGIN, etc., "
         "which prevents routing breaks when changing URLs and provides a single configuration point for all frontend navigation."),

        ("Q5: How is user authentication state shared across components on the React frontend?",
         "We implemented React Context API through AuthContext.jsx. The context maintains the authenticated user's state, "
         "role (admin, employer, or candidate), loading states, and helper methods (login, logout). "
         "This context wraps the root React application, allowing child components to tap into authentication states "
         "via a custom hook (useAuth) without passing props through multiple levels (prop drilling)."),

        ("Q6: How did you make the user interface responsive and mobile-friendly?",
         "The entire UI is built styling-first using Tailwind CSS, adhering to utility-first principles. "
         "We utilize Tailwind's mobile-first breakpoints (such as sm, md, lg, xl) to adjust layouts. "
         "For example, grid layouts use class sets like 'grid grid-cols-1 md:grid-cols-3' to stack content vertically on mobile phones "
         "while automatically displaying three columns side-by-side on desktop displays. We also use flexboxes and viewport-relative units."),

        ("Q7: Explain the data binding and live preview logic in the Live Resume Builder.",
         "We bound the form inputs on the left side of the screen to a React state object representing the resume (containing education, experience, skills, etc.). "
         "As the user types, the input's onChange handler fires, triggering a state update. "
         "The right-hand side preview pane acts as a pure presenter component; it receives the state object as props and renders it live on a CSS-styled print sheet. "
         "When the user clicks 'Save', the React component fires an HTTP request containing the stringified JSON payload, committing it to the database."),

        # PART C
        ("Q8: How does the MVC pattern apply to your Express backend?",
         "We structured the backend using a loose Model-View-Controller (MVC) pattern. "
         "The Model layer represents our database collection schemas defined using Mongoose (e.g. userModel, jobModel). "
         "The View layer is represented by JSON payloads returned by REST APIs (rather than standard HTML pages, since React handles UI). "
         "The Controller layer resides in src/controllers (e.g., jobController.js) and contains isolated async business logic handlers. "
         "Routes in src/routes mapping HTTP methods and paths call these controller functions, keeping routers clean."),

        ("Q9: How is form validation handled on the backend? Why validate on both frontend and backend?",
         "We use the 'express-validator' library to define validation rules as middleware blocks directly on the routes (e.g., verifying emails, password strength, and non-empty text). "
         "We validate on both sides because: frontend validation is for User Experience (UX)—providing instant warnings without network latency; "
         "backend validation is for Security—since malicious users can bypass frontend logic entirely using tools like Postman or curl, "
         "backend validation acts as our line of defense before database inserts."),

        ("Q10: What is the role of your global error handler middleware in Express?",
         "We use a centralized error handler middleware (errorHandler.js) at the bottom of the Express middleware pipeline. "
         "Instead of writing catch blocks that manually format and send error responses everywhere, we catch errors in async controllers "
         "and pass them to next(err). The global error middleware captures all thrown errors, checks if they are operational or database errors, "
         "sets the appropriate HTTP status code (like 400, 401, 500), and sends a clean JSON payload. This prevents sensitive db stack traces from leaking to clients."),

        ("Q11: How does the CareerBot chatbot stream responses in real-time? Explain Server-Sent Events (SSE).",
         "The CareerBot uses Server-Sent Events (SSE) via the /api/chat endpoint. "
         "SSE is an HTTP standard that allows a web server to push real-time updates to a client over a single, long-lived HTTP connection. "
         "The backend sets headers like 'Content-Type: text/event-stream' and 'Cache-Control: no-cache', keeping the response open. "
         "As Gemini AI generates text chunks, the backend writes the data using res.write() in SSE event formats. The React client "
         "reads these chunks stream-by-stream, rendering text word-by-word like ChatGPT."),

        ("Q12: Why did you choose Server-Sent Events (SSE) over WebSockets for the chatbot?",
         "WebSockets are designed for full-duplex, bidirectional communication (e.g. multiplayer games or real-time editing tools). "
         "However, a chatbot only requires unidirectional streaming—the user sends a single query, and the AI streams back a long response. "
         "SSE is lightweight, runs over standard HTTP/HTTPS protocols (avoiding firewall blockages), has built-in auto-reconnection support, "
         "and consumes significantly fewer server resources compared to maintaining full WebSockets protocols."),

        # PART D
        ("Q13: How are your MongoDB collections structured? What relations exist?",
         "We have four main collections defined via Mongoose: "
         "1. User: Stores account credentials, hashed passwords, roles (candidate/employer/admin), and profile settings. "
         "2. Job: Stores title, company name, skills, location, description, and salary. It has a relational reference field (ref: 'User') matching the employer who posted it. "
         "3. Resume: Stores student details, education arrays, experience details, projects, and skills. Ref: 'User' links it to a candidate. "
         "4. InterviewSession: Logs mock interviews. It contains candidate ID, job role, generated questions, submitted answers, and evaluations."),

        ("Q14: What are text indexes in MongoDB, and how did you use them for the search portals?",
         "Text indexes support content searches on string fields in MongoDB. "
         "We created a compound text index on the Job collection matching fields 'title', 'company', and 'skills', "
         "and on the Resume collection matching 'fullName', 'title', and 'skills'. "
         "Using Mongoose's `$text` and `$search` operators, recruiters can search resumes for keywords like 'React' or 'Node' "
         "and retrieve matching candidates sorted by text match score relevance, without requiring expensive regex scans."),

        # PART E
        ("Q15: Why should passwords never be stored in plain text, and how does your project secure them?",
         "Storing plain text passwords leaves user credentials vulnerable in the event of a database breach. "
         "We use 'bcryptjs' to salt and hash passwords before saving them. "
         "Salting adds a random string to the password to protect against rainbow table attacks. "
         "Hashing is a one-way cryptographic function: when logging in, the candidate submits their password, "
         "and bcrypt.compare() hashes it and compares it with the saved hash. The plain text password is never saved or logged."),

        ("Q16: Explain JWT and compare storing tokens in HTTP-only cookies vs. LocalStorage.",
         "JSON Web Token (JWT) is an open standard (RFC 7519) that defines a compact and self-contained way to securely transmit information between parties as a JSON object. "
         "If stored in LocalStorage, JWTs are vulnerable to Cross-Site Scripting (XSS) attacks because any malicious script running in the browser can read local storage. "
         "If stored in HttpOnly cookies, the browser automatically sends the token with every request, but client-side JavaScript cannot read it, preventing token theft. "
         "The drawback is vulnerability to CSRF, which we mitigate by checking request origins in CORS configurations."),

        ("Q17: What is MongoDB query injection, and how does your project defend against it?",
         "Query injection occurs when a user input contains MongoDB query selectors (like `$gt` or `$ne`) which trick the database into bypasses. "
         "For example, sending `{\"username\": {\"$gt\": \"\"}, \"password\": {\"$gt\": \"\"}}` in a login body could return true if unmitigated, logging the attacker in. "
         "We defend against this using the 'express-mongo-sanitize' middleware. "
         "It intercepts all incoming requests (req.body, req.query, req.params) and strips out any keys starting with a dollar sign ($) or containing dots (.), "
         "rendering injection payloads harmless."),

        ("Q18: What does 'xss-clean' middleware do, and how does it protect your application?",
         "XSS (Cross-Site Scripting) clean is a security middleware that sanitizes user inputs from req.body, req.query, and req.params. "
         "It filters out malicious HTML tags, inline scripts, and onload attributes (e.g. `<script>stealCookies()</script>`). "
         "By converting these script characters into safe HTML entities or removing them entirely, it ensures that if a student inputs "
         "malicious code in the resume builder, it will not run inside another user's (or recruiter's) browser when viewed."),

        ("Q19: Explain the rate limiter in your Express server. Why is it used?",
         "We use the 'express-rate-limit' middleware on all api routes (e.g. limit to 300 requests per 15 minutes). "
         "It counts incoming HTTP requests based on the client's IP address. "
         "It protects our backend from Denial of Service (DoS) attacks, brute-force login attempts (by blocking IPs that repeatedly fail login), "
         "and controls costs from our paid third-party AI APIs (Google Gemini) by preventing automated script abuse."),

        # PART F
        ("Q20: What is the role of the '@google/generative-ai' SDK in this project?",
         "The Google Generative AI Node SDK allows our backend API to communicate directly with Google Gemini models. "
         "We use it in two features: "
         "1. CareerBot: An AI chat advisor that answers career-related questions. "
         "2. AI Mock Interview Simulator: Generates position-specific interview questions and evaluates user answers, "
         "returning detailed grading reports."),

        ("Q21: How do you guarantee that Gemini AI returns structured JSON output instead of plain text?",
         "We configure the SDK client parameters. We pass custom system instructions telling the model to return a structured JSON response "
         "matching a predefined schema. In our call, we set the response MIME type to `application/json` and "
         "clearly specify the JSON schema keys (e.g., score, feedback, modelAnswer) in the prompt, forcing the API "
         "to output valid JSON that our backend parses safely using JSON.parse()."),

        ("Q22: How does the backend handle latency or API connection failures with the Gemini AI service?",
         "Third-party AI API calls introduce network latency and could fail. "
         "We handle this by wrapping the API calls in try-catch blocks. If the Gemini API fails, our catch block "
         "gracefully intercepts the error, logs it, and returns a user-friendly API error response (such as a 503 Service Unavailable) "
         "rather than letting the server crash. We also configure connection timeouts to prevent requests from hanging indefinitely."),

        ("Q23: How do you verify database connection security and state on startup?",
         "In server.js, we call a DB connection function before starting the server. "
         "Using Mongoose's mongoose.connect(), we connect to MongoDB Atlas. We listen to event handlers: "
         "'connected' (logs successful DB connection), 'error' (logs connection errors), and 'disconnected'. "
         "If the database connection fails on startup, the server exits loudly using process.exit(1), preventing a half-broken state."),

        ("Q24: How was the project populated with mock data for testing?",
         "We created a seed script (scripts/seedData.js) in the backend folder. "
         "This script uses Mongoose to connect to the DB, clears existing collections (using deleteMany()), "
         "hashes candidate and recruiter passwords using bcrypt, and inserts pre-configured candidates, employers, jobs, "
         "and admin records. This ensures all developers work with identical, clean data during development and testing."),

        ("Q25: If this project were deployed to production, what configurations would you change?",
         "For production deployment: "
         "1. Update the frontend .env.production to reference the live API URL instead of localhost:5001. "
         "2. Change backend .env configuration: NODE_ENV set to 'production', secure cookies enabled, and rate limiter tightened. "
         "3. Store environment variables (mongo connection string, Gemini API key, JWT Secret) securely in the cloud provider dashboard "
         "(like Vercel, Render, or AWS Secrets Manager) instead of committing .env files to GitHub. "
         "4. Build the React app into a static folder ('npm run build') and serve the build assets directly from the backend server to optimize costs."),

        ("Q26: What was the most challenging feature in this project, and how did you resolve it?",
         "The most challenging feature was implementing the AI Mock Interview Simulator. "
         "It required linking dynamic prompt generation, input handling on the UI, and parsing JSON responses from the Gemini API. "
         "Since AI models are non-deterministic, they occasionally outputted markdown backticks (like ```json ... ```) "
         "which caused JSON.parse() to throw errors. We solved this by setting the response format to JSON in Gemini's SDK config "
         "and writing a robust regex utility on the backend to strip any accidental markdown formatting before parsing the response.")
    ]

    for q, a in questions:
        add_custom_heading(q, level=3, space_before=12, space_after=4)
        p_ans = doc.add_paragraph()
        p_ans.add_run("Answer: ").bold = True
        p_ans.add_run(a)
        # Separator spacing
        p_sep = doc.add_paragraph()
        p_sep.paragraph_format.space_after = Pt(6)

    # Save the document
    output_filename = "CareerHub_Viva_Preparation_Guide.docx"
    doc.save(output_filename)
    print(f"Success! Saved document to {output_filename}")

if __name__ == "__main__":
    build_docx()
